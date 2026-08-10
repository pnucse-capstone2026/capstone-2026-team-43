"""Sentry Host Agent 중간 보고서 Word 문서 생성."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "Sentry_Host_Agent_중간보고서.docx"


def set_doc_font(doc: Document, name: str = "맑은 고딕") -> None:
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_title_page(doc: Document) -> None:
    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Sentry Host Agent\n개발 중간 보고서")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"작성일: {date.today().strftime('%Y년 %m월 %d일')}\n").font.size = Pt(12)
    meta.add_run("프로젝트: host-based DLP (Data Loss Prevention) 에이전트\n").font.size = Pt(12)
    meta.add_run("대상 OS: Windows 10 / 11").font.size = Pt(12)

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
    doc.add_paragraph()


def build_report() -> Document:
    doc = Document()
    set_doc_font(doc)

    # 여백
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    add_title_page(doc)

    # ── 1. 개요 ──────────────────────────────────────────────────────────────
    add_heading(doc, "1. 개요")
    add_para(
        doc,
        "본 보고서는 Windows 호스트 기반 DLP(Data Loss Prevention) 솔루션인 "
        "「Sentry Host Agent」의 개발 진행 상황을 중간 점검 형식으로 정리한 문서이다. "
        "에이전트는 클립보드, 이메일, 웹 업로드, USB 및 네트워크 공유 등 "
        "다양한 정보 반출 경로를 모니터링하고, 정규표현식 1차 필터링과 "
        "AI 서버 연동(현재 mock 모드)을 통해 민감 정보 유출을 차단하는 것을 목표로 한다.",
    )

    add_heading(doc, "1.1 프로젝트 목적", 2)
    add_bullets(doc, [
        "엔드포인트(호스트)에서 민감 정보의 외부 반출을 사전에 탐지·차단",
        "사용자 업무 효율을 해치지 않도록 반출 경로·목적지 기반의 선택적 검사 적용",
        "정규표현식 1차 게이트 + AI 서버 2차 판단의 2단계 분석 파이프라인 구축",
        "탐지 이벤트의 로컬 저장(JSONL) 및 향후 웹 대시보드 연동 기반 마련",
    ])

    add_heading(doc, "1.2 개발 범위", 2)
    add_table(doc,
        ["구분", "내용", "상태"],
        [
            ["1단계", "클립보드 제어 (Python)", "완료"],
            ["2단계", "네트워크/메일 반출 탐지 (Python)", "진행 중"],
            ["3단계", "USB·네트워크 공유 (미니필터 C/C++)", "진행 중"],
            ["공통", "이벤트 로깅, AI 페이로드 정제, Windows 서비스", "완료"],
        ],
    )

    # ── 2. 시스템 아키텍처 ───────────────────────────────────────────────────
    add_heading(doc, "2. 시스템 아키텍처")

    add_heading(doc, "2.1 전체 구조", 2)
    add_para(doc, "프로젝트는 모듈형 단일 레포지토리(sentry-host-agent)로 구성되며, "
             "각 반출 채널별 모듈과 공통 통신·로깅 모듈로 분리되어 있다.")
    add_table(doc,
        ["디렉터리", "역할"],
        [
            ["clipboard_ctrl/", "클립보드 붙여넣기 인터셉트 및 정책 적용"],
            ["network_hook/", "Outlook, SMTP, HTTP, 파일 내용 추출"],
            ["usb_minifilter/", "미니필터 커널 드라이버 + Python 유저모드 클라이언트"],
            ["core_comm/", "AI API 클라이언트, 페이로드 정제, 이벤트 로깅"],
            ["config/", "settings.yaml, regex_patterns.json, channel_policy.json 등"],
            ["main_agent.py", "에이전트 진입점 (mode: all / system / user)"],
            ["sentry_service.py", "Windows 서비스 (Session 0 채널)"],
            ["sentry_user_agent.py", "유저 세션 에이전트 (Task Scheduler)"],
        ],
    )

    add_heading(doc, "2.2 분석 파이프라인", 2)
    add_para(doc, "모든 채널은 동일한 DLP 파이프라인을 공유한다.")
    add_bullets(doc, [
        "1차 게이트: RuleFilter — regex_patterns.json 기반 정규표현식 매칭",
        "데이터 정제: PayloadBuilder — 매칭 구간 스니펫 추출, PII 부분 마스킹, 메타데이터 구성",
        "2차 판단: ApiClient — AI 서버 POST (현재 ai_base_url 미설정 시 mock 모드)",
        "조치: block → 차단 + 알림 + 로그 / review → 허용 + 로그 / allow → 통과",
    ])

    add_heading(doc, "2.3 Windows 서비스 이중 구조", 2)
    add_para(doc,
        "Windows Session 0 격리로 인해 클립보드·Outlook·UI 알림은 사용자 세션에서만 동작 가능하다. "
        "이에 따라 두 컴포넌트로 분리하였다.",
    )
    add_table(doc,
        ["컴포넌트", "실행 방식", "담당 채널"],
        [
            ["SentryDLPService", "Windows 서비스 (자동 시작, 크래시 재시작)", "SMTP, HTTP, FileGuard"],
            ["SentryUserAgent", "Task Scheduler (로그온 트리거)", "클립보드, Outlook, 알림 팝업"],
        ],
    )

    # ── 3. 구현 완료 항목 ─────────────────────────────────────────────────────
    add_heading(doc, "3. 구현 완료 항목")

    add_heading(doc, "3.1 클립보드 제어 (clipboard_ctrl)", 2)
    add_bullets(doc, [
        "WH_KEYBOARD_LL 글로벌 키보드 훅으로 Ctrl+V 붙여넣기 사전 인터셉트",
        "AddClipboardFormatListener + WM_CLIPBOARDUPDATE로 클립보드 변경 추적",
        "TextExtractor: CF_UNICODETEXT, CF_HTML 형식 텍스트 추출 및 정규화",
        "PasteInspector: paste_policy.json 기반 allowlist / inspect_list / default_action",
        "외부 반출 가능 앱(Slack, Chrome, Discord 등)에 붙여넣을 때만 검사 — 내부 앱(Excel, Word)은 허용",
        "RuleFilter: 18종 민감 정보 패턴 (주민번호, 카드번호, API Key 등)",
        "Notifier: Tkinter 기반 화면 중앙 최상위 팝업 알림",
    ])

    add_heading(doc, "3.2 공통 모듈 (core_comm)", 2)
    add_bullets(doc, [
        "LocalEventStore: logs/events.jsonl 스레드 안전 JSONL 저장, 10MB 로테이션",
        "EventLogger: 로컬 저장 + 선택적 대시보드 즉시 전송(send_immediately)",
        "PayloadBuilder: 매칭 ±150자 스니펫, 패턴별 부분 마스킹, severity 정렬",
        "ApiClient: mock 모드(critical/high→block, medium→review) 및 실제 AI 서버 연동 준비",
    ])

    add_heading(doc, "3.3 운영 배포 (install/)", 2)
    add_bullets(doc, [
        "setup_agent.ps1: Windows 서비스 등록 + Task Scheduler 작업 등록 + 크래시 재시작 정책",
        "remove_agent.ps1: 서비스·작업·프로세스 일괄 제거",
        "main_agent.py run(mode): all(개발), system(서비스), user(유저 에이전트) 분리",
    ])

    # ── 4. 주요 기술 결정 ───────────────────────────────────────────────────
    add_heading(doc, "4. 주요 기술 결정")

    add_table(doc,
        ["항목", "결정", "근거"],
        [
            ["클립보드 감지", "Ctrl+V 인터셉트 + 목적지 정책", "내부 앱 작업 효율 유지, 외부 앱만 검사"],
            ["클립보드 파일/이미지", "클립보드에서 제외", "지연·추출 비용 대비 반출 경로 검사가 효과적"],
            ["AI 분석", "정규식 hit 시에만 AI 호출", "전체 텍스트 AI 전송 비용·지연 방지"],
            ["AI 페이로드", "스니펫 + 마스킹", "토큰 절감 및 PII 최소화"],
            ["미니필터 Fail-Open", "오류 시 허용", "업무 중단 최소화 (fail_open 설정으로 전환 가능)"],
            ["서비스 구조", "2-Process (Session 0 + User)", "Windows Session 0 격리 제약"],
        ],
    )

    # ── 5. 설정 파일 ─────────────────────────────────────────────────────────
    add_heading(doc, "5. 주요 설정 파일")
    add_table(doc,
        ["파일", "용도"],
        [
            ["config/settings.yaml", "AI 서버 URL, 로그 경로, 임계값, 로그 레벨"],
            ["config/regex_patterns.json", "18종 민감 정보 정규표현식 (severity: critical/high/medium)"],
            ["config/paste_policy.json", "붙여넣기 inspect_list / allowlist / default_action"],
            ["config/channel_policy.json", "Outlook, SMTP, HTTP, file_guard 채널 on/off"],
        ],
    )

    # ── 6. 테스트 현황 ───────────────────────────────────────────────────────
    add_heading(doc, "6. 테스트 현황")
    add_para(doc, "pytest 기반 단위 테스트 54건 작성, 54건 통과.")
    add_para(doc, "※ 2·3단계(네트워크/메일, USB·네트워크 공유) 관련 테스트는 개발 완료 후 별도 반영 예정.")
    add_table(doc,
        ["테스트 파일", "대상", "결과"],
        [
            ["tests/test_clipboard.py", "TextExtractor, RuleFilter, PasteInspector", "19 passed"],
            ["tests/test_event_logger.py", "LocalEventStore, EventLogger", "11 passed"],
            ["tests/test_payload_builder.py", "PayloadBuilder, ApiClient mock, 마스킹", "24 passed"],
        ],
    )

    # ── 7. 미완료·향후 과제 ─────────────────────────────────────────────────
    add_heading(doc, "7. 미완료 및 향후 과제")

    add_heading(doc, "7.1 2단계 — 네트워크/메일 반출 탐지 (진행 중)", 2)
    add_bullets(doc, [
        "OutlookHook: win32com ItemSend 이벤트 — 외부 수신자 메일 본문·첨부 검사, 차단 시 Drafts 이동",
        "SmtpProxy: aiosmtpd 기반 로컬 SMTP 프록시(기본 2525) — Thunderbird 등 비-Outlook 클라이언트 대응",
        "HttpMonitor: WinDivert/pydivert 기반 HTTP 업로드 감시 skeleton → 본 구현 및 HTTPS 한계 문서화",
        "FileInspector: docx, pdf, xlsx, txt, csv 등 파일 텍스트 추출 (샘플링 제한 적용)",
        "channel_policy.json 채널별 on/off 및 internal_domains 설정",
        "FTP 탐지 (로컬 FTP 프록시 또는 WFP 후킹)",
    ])

    add_heading(doc, "7.2 3단계 — USB·네트워크 공유 (진행 중)", 2)
    add_bullets(doc, [
        "sentry_filter.sys: Filter Manager 미니필터 (altitude 265000) — WDK 빌드 및 EV 코드 서명",
        "IRP_MJ_CREATE / WRITE / SET_INFORMATION pre-op — USB·SMB 볼륨 쓰기·이동 감시",
        "FltSendMessage ↔ FilterGetMessage — 커널↔유저모드 동기 통신",
        "file_guard.py: fltlib.dll FilterConnectCommunicationPort 클라이언트",
        "Fail-Open 정책: 에이전트 미연결·타임아웃(5초) 시 허용",
        "install.ps1 / uninstall.ps1: fltMC load/unload 기반 드라이버 설치·현장 배포 검증",
        "클라우드 동기 폴더(OneDrive/Dropbox) 경로 감시",
    ])

    add_heading(doc, "7.3 단기 (운영 안정성)", 2)
    add_bullets(doc, [
        "채널 워치독: 개별 채널 크래시 시 자동 재시작",
        "설정 hot-reload: settings.yaml / channel_policy.json 변경 시 재시작 없이 반영",
        "이벤트 배치 전송: EventLogger.flush_to_server() 주기 실행",
        "에이전트 heartbeat: 대시보드에 상태·버전 주기 보고",
    ])

    add_heading(doc, "7.4 장기 (연동·배포)", 2)
    add_bullets(doc, [
        "웹 대시보드 연동 (이벤트 수집 API, 정책 원격 동기화)",
        "AI 서버 실연동 (ai_base_url 설정, medium severity AI 위임)",
        "MSI/NSIS 인스톨러 (Python + sentry_filter.sys 일괄 설치)",
        "GPO/ADMX 그룹 정책 템플릿",
        "VM 드래그앤드롭·VM 내부 반출 탐지",
    ])

    # ── 8. 결론 ──────────────────────────────────────────────────────────────
    add_heading(doc, "8. 결론")
    add_para(
        doc,
        "Sentry Host Agent는 1단계 클립보드 제어와 공통 모듈(이벤트 로깅, AI 페이로드 정제, "
        "Windows 서비스) 구현을 완료하였다. "
        "2단계(네트워크/메일 반출 탐지)와 3단계(USB·네트워크 공유 미니필터)는 "
        "기본 골격 및 skeleton 수준까지 진행 중이며, 향후 과제로 본 구현·현장 검증을 진행할 예정이다.\n\n"
        "다음 단계로는 2·3단계 채널의 완성, 미니필터 드라이버 빌드·서명·현장 테스트, "
        "웹 대시보드·AI 서버와의 연동을 순차적으로 추진할 계획이다.",
    )

    # ── 부록: 파일 목록 ──────────────────────────────────────────────────────
    doc.add_page_break()
    add_heading(doc, "부록 A. 주요 소스 파일 목록")
    files = [
        "main_agent.py", "sentry_service.py", "sentry_user_agent.py",
        "clipboard_ctrl/clipboard_hook.py", "clipboard_ctrl/paste_inspector.py",
        "clipboard_ctrl/rule_filter.py", "clipboard_ctrl/text_extractor.py", "clipboard_ctrl/notifier.py",
        "network_hook/outlook_hook.py", "network_hook/smtp_proxy.py", "network_hook/http_monitor.py",
        "network_hook/file_inspector.py",
        "core_comm/payload_builder.py", "core_comm/api_client.py",
        "core_comm/event_logger.py", "core_comm/local_store.py",
        "usb_minifilter/driver/sentry_filter.c", "usb_minifilter/driver/sentry_comm.h",
        "usb_minifilter/file_guard.py",
        "install/setup_agent.ps1", "usb_minifilter/install/install.ps1",
    ]
    for f in files:
        doc.add_paragraph(f, style="List Bullet")

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_report()
    doc.save(str(OUTPUT))
    print(f"보고서 생성 완료: {OUTPUT}")


if __name__ == "__main__":
    main()
