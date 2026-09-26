"""Shared utilities for channel benchmark scripts."""

from __future__ import annotations

import csv
import logging
import os
import time
from email import message_from_bytes
from pathlib import Path
from statistics import mean
from typing import Any, Optional

import yaml

PROJECT_ROOT = Path(__file__).resolve().parent.parent
BENCH_DIR    = PROJECT_ROOT / "tests" / "bench"
RESULTS_DIR  = PROJECT_ROOT / "tests" / "results"
SETTINGS     = PROJECT_ROOT / "config" / "settings.yaml"
PATTERNS     = PROJECT_ROOT / "config" / "regex_patterns.json"


def default_output_path(channel_slug: str, fmt: str = "xlsx") -> str:
    """채널별 기본 결과 파일 경로 (tests/results/bench_{channel}.{fmt})."""
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    return str(RESULTS_DIR / f"bench_{channel_slug}.{fmt}")

DEFAULT_EXTS  = None   # None = bench/ 내 전체 확장자
DEFAULT_SIZES = None   # None = bench/ 내 전체 크기 구간

# 크기 구간 정렬 순서 (bench 폴더에 50mb까지 존재; 100mb 없음)
SIZE_ORDER = {
    "10kb": 0, "100kb": 1, "1mb": 2, "10mb": 3, "50mb": 4, "100mb": 5,
}

# 확장자 → 보고서용 한글 유형명
EXT_LABELS: dict[str, str] = {
    "txt": "텍스트", "csv": "CSV", "log": "로그", "md": "Markdown",
    "docx": "Word", "pdf": "PDF", "xlsx": "Excel", "pptx": "PowerPoint",
    "hwp": "한글(HWP)", "hwpx": "한글(HWPX)", "rtf": "RTF",
    "zip": "ZIP",
    "py": "Python", "js": "JavaScript", "ts": "TypeScript",
    "tsx": "TypeScript", "jsx": "JavaScript", "java": "Java",
    "cs": "C#", "go": "Go", "c": "C", "cpp": "C++", "h": "Header",
    "rb": "Ruby", "php": "PHP", "sh": "Shell", "bat": "Batch",
    "ps1": "PowerShell", "sql": "SQL",
    "htm": "HTML", "html": "HTML", "xml": "XML", "json": "JSON",
    "yaml": "YAML", "yml": "YAML", "env": "ENV",
}

# 내부 집계용 필드
CSV_FIELDS = [
    "type_label", "ext", "size",
    "channel", "extract", "regex", "send", "ai", "recv", "block", "total", "n_runs",
]

# Excel 보고서 컬럼 (유형/용량별 1행)
REPORT_HEADERS = [
    "파일 유형", "파일 크기",
    "extract", "regex", "send", "ai", "recv", "block", "total",
]


def ext_label(ext: str) -> str:
    return EXT_LABELS.get(ext.lower(), ext.upper())


def discover_sizes() -> list[str]:
    """bench/ 하위 크기 디렉터리 목록 (정렬됨)."""
    sizes = [
        p.name for p in BENCH_DIR.iterdir()
        if p.is_dir() and list(p.glob("sample.*"))
    ]
    return sorted(sizes, key=lambda s: SIZE_ORDER.get(s, 99))


def load_settings() -> dict[str, Any]:
    with SETTINGS.open(encoding="utf-8") as f:
        return yaml.safe_load(f)


def build_clients(mock: bool) -> tuple[Any, Any, Any, Any]:
    """Return (api, payload_builder, rule_filter, file_inspector)."""
    import sys
    sys.path.insert(0, str(PROJECT_ROOT))

    from clipboard_ctrl.rule_filter import RuleFilter
    from core_comm.api_client import ApiClient
    from core_comm.payload_builder import PayloadBuilder
    from network_hook.file_inspector import FileInspector

    settings = load_settings()
    server   = settings.get("server", {})
    agent    = settings.get("agent", {})

    ai_url = "" if mock else server.get("ai_base_url", "")
    api = ApiClient(base_url=ai_url, timeout=server.get("request_timeout_sec", 10))
    pb  = PayloadBuilder(
        max_text_chars=agent.get("ai_max_text_chars", 0),
        max_per_pattern=agent.get("ai_max_matches_per_pattern", 10),
        context_chars=agent.get("ai_context_chars", 200),
        user_id=os.environ.get("USERNAME", "bench"),
    )
    return api, pb, RuleFilter(PATTERNS), FileInspector()


def collect_files(
    exts: Optional[list[str]] = None,
    sizes: Optional[list[str]] = None,
) -> list[Path]:
    """bench/ 내 sample.* 파일 전체 또는 필터링 subset."""
    use_sizes = sizes if sizes else discover_sizes()
    files: list[Path] = []

    for size_dir in sorted(BENCH_DIR.iterdir(), key=lambda p: SIZE_ORDER.get(p.name, 99)):
        if not size_dir.is_dir() or size_dir.name not in use_sizes:
            continue
        for f in sorted(size_dir.glob("sample.*")):
            ext = f.suffix.lstrip(".").lower()
            if exts and ext not in exts:
                continue
            files.append(f)

    return files


def run_dlp_core(
    text: str,
    rf: Any,
    api: Any,
    pb: Any,
    channel_name: str,
    process_name: str,
) -> dict[str, float]:
    """Regex + AI on extracted text."""
    t0 = time.perf_counter()
    hits = rf.match(text)
    t_regex = (time.perf_counter() - t0) * 1000

    bench = {"regex": t_regex, "send": 0.0, "ai": 0.0, "recv": 0.0}
    if hits:
        payload = pb.build(text, hits, channel_name, process_name)
        result  = api.analyze(payload)
        bench["send"] = result.bench.t_send_ms
        bench["ai"]   = result.bench.t_ai_ms
        bench["recv"] = result.bench.t_recv_ms
    return bench


def aggregate(
    runs: list[dict[str, float]],
    block_ms: float,
) -> dict[str, float]:
    """Drop min(core+channel) run when N>=5, average the rest."""
    if not runs:
        return {}

    def _total(r: dict[str, float]) -> float:
        return (
            r.get("channel", 0) + r["extract"] + r["regex"]
            + r["send"] + r["ai"] + r["recv"]
        )

    if len(runs) >= 5:
        drop_idx = min(enumerate(runs), key=lambda x: _total(x[1]))[0]
        kept = [r for i, r in enumerate(runs) if i != drop_idx]
    else:
        kept = runs

    out = {
        "channel": round(mean(r.get("channel", 0) for r in kept), 2),
        "extract": round(mean(r["extract"] for r in kept), 2),
        "regex":   round(mean(r["regex"]   for r in kept), 2),
        "send":    round(mean(r["send"]    for r in kept), 2),
        "ai":      round(mean(r["ai"]      for r in kept), 2),
        "recv":    round(mean(r["recv"]    for r in kept), 2),
        "block":   round(block_ms, 2),
    }
    out["total"] = round(
        out["channel"] + out["extract"] + out["regex"]
        + out["send"] + out["ai"] + out["recv"] + out["block"],
        2,
    )
    out["n_runs"] = len(kept)
    return out


def warmup_ai(api: Any, pb: Any, rf: Any, fi: Any) -> None:
    sample = BENCH_DIR / "10kb" / "sample.txt"
    if not sample.exists():
        return
    text = fi.extract_from_path(sample) or ""
    hits = rf.match(text)
    if hits:
        api.analyze(pb.build(text, hits, "bench", "warmup"))
        print("[warmup] AI cold-start done")


def row_display_type(row: dict[str, Any]) -> str:
    """보고서 '파일 유형' 표시 — ts/tsx 등 동일 한글명 확장자 구분."""
    ext = row.get("ext", "")
    label = row.get("type_label", ext.upper())
    return f"{label} (.{ext})" if ext else label


def save_excel_report(rows: list[dict[str, Any]], output_path: str, sheet_name: str = "results") -> None:
    """유형/용량별 행 + 병합 셀 형식 Excel 보고서."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    out = PROJECT_ROOT / (output_path if output_path.lower().endswith(".xlsx") else output_path + ".xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)

    # 확장자 → 크기 순 정렬 (동일 한글 유형명 ts/tsx 등은 .ext 로 구분)
    sorted_rows = sorted(
        rows,
        key=lambda r: (r.get("ext", ""), SIZE_ORDER.get(r.get("size", ""), 99)),
    )

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, name in enumerate(REPORT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for row_idx, row in enumerate(sorted_rows, start=2):
        # extract = 채널 오버헤드 + 파일 추출 (보고서 단일 컬럼)
        extract_val = round(row.get("channel", 0) + row.get("extract", 0), 2)
        values = [
            row_display_type(row),
            row.get("size", ""),
            extract_val,
            row.get("regex", 0),
            row.get("send", 0),
            row.get("ai", 0),
            row.get("recv", 0),
            row.get("block", 0),
            row.get("total", 0),
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = border
            if col >= 3:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "0.00"

    # 파일 유형 열 병합 (연속 동일 표시명 — 확장자별 5행)
    if sorted_rows:
        merge_start = 2
        prev_label = row_display_type(sorted_rows[0])
        for i, row in enumerate(sorted_rows[1:], start=3):
            label = row_display_type(row)
            if label != prev_label:
                if i - 1 > merge_start:
                    ws.merge_cells(start_row=merge_start, start_column=1,
                                   end_row=i - 1, end_column=1)
                    ws.cell(row=merge_start, column=1).alignment = Alignment(vertical="center")
                merge_start = i
                prev_label = label
        last_row = len(sorted_rows) + 1
        if last_row > merge_start:
            ws.merge_cells(start_row=merge_start, start_column=1,
                           end_row=last_row, end_column=1)
            ws.cell(row=merge_start, column=1).alignment = Alignment(vertical="center")

    widths = [14, 10, 10, 8, 10, 8, 10, 8, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(out)
    print(f"\n[saved] {out.resolve()}  ({len(sorted_rows)} rows)")


def save_excel_raw(rows: list[dict[str, Any]], output_path: str, sheet_name: str = "raw") -> None:
    """상세 raw 데이터 시트 (선택)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    out = PROJECT_ROOT / output_path
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    fields = CSV_FIELDS
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for col, name in enumerate(fields, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.fill = header_fill
        c.font = Font(bold=True)
    for ri, row in enumerate(rows, start=2):
        for ci, key in enumerate(fields, start=1):
            ws.cell(row=ri, column=ci, value=row.get(key, ""))
    wb.save(out)


def save_results(rows: list[dict[str, Any]], output_path: str, sheet_name: str = "results") -> None:
    """Save formatted Excel report (.xlsx) or flat CSV."""
    if output_path.lower().endswith(".csv"):
        out = PROJECT_ROOT / output_path
        out.parent.mkdir(parents=True, exist_ok=True)
        flat_fields = ["type_label", "size", "ext"] + REPORT_HEADERS[2:]
        with out.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=flat_fields, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                flat = {
                    "type_label": row_display_type(row),
                    "size": row.get("size"),
                    "ext": row.get("ext"),
                    "extract": round(row.get("channel", 0) + row.get("extract", 0), 2),
                    "regex": row.get("regex"),
                    "send": row.get("send"),
                    "ai": row.get("ai"),
                    "recv": row.get("recv"),
                    "block": row.get("block"),
                    "total": row.get("total"),
                }
                writer.writerow(flat)
        print(f"\n[saved] {out.resolve()}")
    else:
        save_excel_report(rows, output_path, sheet_name)


def print_row(type_label: str, ext: str, size: str, agg: dict[str, float]) -> None:
    extract_show = round(agg.get("channel", 0) + agg.get("extract", 0), 2)
    print(
        f"  {type_label:12s} ({ext:5s})  {size:6s}  "
        f"extract={extract_show:8.2f}  regex={agg['regex']:6.2f}  "
        f"send={agg['send']:8.2f}  ai={agg['ai']:6.2f}  "
        f"recv={agg['recv']:8.2f}  block={agg['block']:6.2f}  total={agg['total']:8.2f}  "
        f"(n={agg['n_runs']})"
    )


# ── Multipart helpers (mail / drive mitm simulation) ─────────────────────────

def build_mail_multipart(path: Path) -> tuple[bytes, str]:
    """Gmail-style multipart/related body."""
    boundary = "bench_mail_boundary"
    raw      = path.read_bytes()
    filename = path.name
    meta = (
        f'Content-Type: application/json; charset=UTF-8\r\n\r\n'
        f'{{"filename":"{filename}"}}\r\n'
    )
    file_part = (
        f"Content-Type: application/octet-stream\r\n"
        f'Content-Disposition: attachment; filename="{filename}"\r\n\r\n'
    ).encode() + raw + b"\r\n"

    body = (
        f"--{boundary}\r\n".encode()
        + meta.encode()
        + f"--{boundary}\r\n".encode()
        + file_part
        + f"--{boundary}--\r\n".encode()
    )
    ct = f'multipart/related; boundary="{boundary}"'
    return body, ct


def build_drive_multipart(path: Path) -> tuple[bytes, str]:
    """Google Drive web UI-style multipart/related body."""
    boundary = "bench_drive_boundary"
    raw      = path.read_bytes()
    filename = path.name
    mime     = "text/plain" if path.suffix == ".txt" else "application/octet-stream"
    meta = (
        f"Content-Type: application/json; charset=UTF-8\r\n\r\n"
        f'{{"title":"{filename}","mimeType":"{mime}"}}\r\n'
    )
    file_part = (
        f"Content-Type: {mime}\r\n\r\n"
    ).encode() + raw + b"\r\n"

    body = (
        f"--{boundary}\r\n".encode()
        + meta.encode()
        + f"--{boundary}\r\n".encode()
        + file_part
        + f"--{boundary}--\r\n".encode()
    )
    ct = f'multipart/related; boundary="{boundary}"'
    return body, ct


def parse_multipart_for_bench(
    body: bytes,
    ct: str,
    fi: Any,
) -> tuple[float, float, str]:
    """Parse multipart (mitm path) and extract file text.

    Returns (mitm_ms, extract_ms, text).
    mitm_ms  = multipart structure parse + part walk
    extract_ms = FileInspector on file payload
    """
    t0 = time.perf_counter()
    try:
        mime = message_from_bytes(
            b"Content-Type: " + ct.encode() + b"\r\n\r\n" + body
        )
    except Exception:
        return (time.perf_counter() - t0) * 1000, 0.0, ""

    payload: bytes = b""
    filename = "upload.bin"
    inferred = ""

    for part in mime.walk():
        if part.is_multipart():
            continue
        part_payload = part.get_payload(decode=True)
        if not part_payload:
            continue
        pct = (part.get_content_type() or "").lower()
        fname = part.get_filename() or ""

        if "json" in pct:
            try:
                import json
                meta = json.loads(part_payload.decode("utf-8", errors="ignore"))
                if isinstance(meta, dict):
                    inferred = (
                        meta.get("title") or meta.get("name")
                        or meta.get("filename") or meta.get("originalFilename") or inferred
                    )
            except Exception:
                pass
            continue

        payload = part_payload
        filename = fname or inferred or filename

    t_mitm = (time.perf_counter() - t0) * 1000

    if not payload:
        return t_mitm, 0.0, ""

    t0 = time.perf_counter()
    try:
        text = fi.extract_from_bytes(payload, filename) or ""
    except Exception:
        text = ""
    t_extract = (time.perf_counter() - t0) * 1000
    return t_mitm, t_extract, text


def run_bench_loop(
    channel_label: str,
    channel_name: str,
    files: list[Path],
    run_once_fn: Any,
    repeat: int,
    block_ms: float,
    output_path: str,
    sheet_name: str = "results",
) -> None:
    """Generic benchmark loop."""
    print(f"[output] {Path(output_path).resolve()}")
    rows: list[dict[str, Any]] = []
    for path in files:
        ext  = path.suffix.lstrip(".").lower()
        size = path.parent.name
        runs: list[dict[str, float]] = []

        for _ in range(repeat):
            r = run_once_fn(path)
            if r:
                runs.append(r)

        if not runs:
            print(f"  SKIP  {ext:6s}  {size:6s}  (empty)")
            continue

        agg = aggregate(runs, block_ms)
        label = ext_label(ext)
        row = {"type_label": label, "ext": ext, "size": size, **agg}
        rows.append(row)
        print_row(label, ext, size, agg)

    if output_path and rows:
        save_results(rows, output_path, sheet_name)


def make_arg_parser(description: str, channel_slug: str) -> Any:
    import argparse
    default_out = default_output_path(channel_slug)
    p = argparse.ArgumentParser(description=description)
    p.add_argument(
        "--ext", nargs="*", default=None,
        help="확장자 필터 (미지정 = bench/ 전체)",
    )
    p.add_argument(
        "--size", nargs="*", default=None,
        help="크기 구간 필터 (미지정 = bench/ 전체: 10kb~50mb)",
    )
    p.add_argument("--repeat",   type=int,  default=5)
    p.add_argument("--block-ms", type=float, default=90.0)
    p.add_argument(
        "--out", default=default_out,
        help=f"결과 파일 (.xlsx 또는 .csv, 기본: {default_out})",
    )
    p.add_argument("--mock",     action="store_true")
    p.add_argument("--no-warmup", action="store_true")
    return p


def resolve_filters(exts: Optional[list[str]], sizes: Optional[list[str]]) -> tuple[Optional[list[str]], list[str]]:
    """빈 리스트 → None 처리, sizes 미지정 시 bench 전체."""
    exts_out  = exts if exts else None
    sizes_out = sizes if sizes else discover_sizes()
    return exts_out, sizes_out
