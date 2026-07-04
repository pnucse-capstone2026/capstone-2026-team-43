"""network_hook 모듈 단위 테스트."""

import io
import json
import tempfile
from pathlib import Path

import pytest

from network_hook.file_inspector import FileInspector

# ── FileInspector ─────────────────────────────────────────────────────────────

@pytest.fixture
def fi() -> FileInspector:
    return FileInspector(sample_bytes=64 * 1024)


class TestFileInspectorText:
    def test_txt_utf8(self, fi: FileInspector, tmp_path: Path) -> None:
        p = tmp_path / "hello.txt"
        p.write_text("주민번호: 900101-1234567", encoding="utf-8")
        result = fi.extract_from_path(p)
        assert result is not None
        assert "주민번호" in result

    def test_txt_cp949(self, fi: FileInspector, tmp_path: Path) -> None:
        p = tmp_path / "hello.txt"
        p.write_bytes("안녕하세요".encode("cp949"))
        result = fi.extract_from_path(p)
        assert result is not None
        assert "안녕" in result

    def test_csv(self, fi: FileInspector, tmp_path: Path) -> None:
        p = tmp_path / "data.csv"
        p.write_text("name,phone\n홍길동,010-1234-5678", encoding="utf-8")
        result = fi.extract_from_path(p)
        assert result is not None
        assert "010-1234-5678" in result

    def test_unsupported_returns_none_or_str(self, fi: FileInspector, tmp_path: Path) -> None:
        p = tmp_path / "file.bin"
        p.write_bytes(bytes(range(256)))
        # 바이너리라도 None이거나 str 이어야 함
        result = fi.extract_from_path(p)
        assert result is None or isinstance(result, str)

    def test_sample_bytes_limit(self, fi: FileInspector, tmp_path: Path) -> None:
        fi_small = FileInspector(sample_bytes=10)
        p = tmp_path / "big.txt"
        p.write_text("a" * 1000, encoding="utf-8")
        result = fi_small.extract_from_path(p)
        assert result is not None
        assert len(result) <= 10

    def test_nonexistent_path(self, fi: FileInspector) -> None:
        result = fi.extract_from_path(Path("nonexistent_xyz.txt"))
        assert result is None


class TestFileInspectorDocx:
    def test_docx_extract(self, fi: FileInspector) -> None:
        try:
            import docx as _docx
        except ImportError:
            pytest.skip("python-docx 미설치")

        import docx
        doc = docx.Document()
        doc.add_paragraph("주민번호: 900101-1234567")
        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()

        result = fi.extract_from_bytes(data, "secret.docx")
        assert result is not None
        assert "주민번호" in result


class TestFileInspectorXlsx:
    def test_xlsx_extract(self, fi: FileInspector) -> None:
        try:
            import openpyxl as _openpyxl
        except ImportError:
            pytest.skip("openpyxl 미설치")

        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "이름"
        ws["B1"] = "전화번호"
        ws["A2"] = "홍길동"
        ws["B2"] = "010-9876-5432"
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()

        result = fi.extract_from_bytes(data, "contacts.xlsx")
        assert result is not None
        assert "010-9876-5432" in result


class TestFileInspectorPdf:
    def test_pdf_extract(self, fi: FileInspector) -> None:
        try:
            import fitz as _fitz
        except ImportError:
            pytest.skip("PyMuPDF 미설치")

        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "카드번호: 4532-1234-5678-9012")
        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()

        result = fi.extract_from_bytes(data, "report.pdf")
        assert result is not None
        assert "4532" in result
